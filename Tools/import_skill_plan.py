"""Extract a skill workbook into typed authoring data (requires openpyxl).

Usage: python Tools/import_skill_plan.py <workbook.xlsx>
This reads source cells only; game mechanics remain in the C# authoring recipes.
"""
import argparse
import json
import math
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
CONTENT = ROOT / "Assets/_Project/Content"
TRIPLE = r"(\d+(?:\.\d+)?)/(\d+(?:\.\d+)?)/(\d+(?:\.\d+)?)"


def extract(path):
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        # Keep the current plan's stable IDs and evolution graph while updating its numbers.
        baseline = json.loads((CONTENT / "SkillPlan.json").read_text(encoding="utf-8-sig"))
        ids = {row["Dex"]: row["Id"] for row in baseline["Rows"]}
        stats = {}
        for row in workbook["성급별 스탯"].iter_rows(min_row=4, values_only=True):
            if not row[0]:
                continue
            dex = int(row[0])
            if dex in stats:
                raise ValueError(f"Duplicate stats dex: {dex}")
            stats[dex] = row
        rows = []
        for row in workbook["코스트 및 스킬"].iter_rows(min_row=2, values_only=True):
            dex = int(row[0])
            stat = stats[dex]
            if row[1] != stat[1] or row[2] != stat[2]:
                raise ValueError(f"Skill/stats identity mismatch: {dex}")
            if any(not isinstance(v, (int, float)) or not math.isfinite(v) for v in stat[3:15]):
                raise ValueError(f"Missing/invalid cached stats: {dex}. Recalculate the workbook first.")
            effect = row[4]
            damage = [{"Values": list(map(float, match.groups()))} for match in re.finditer(
                TRIPLE + r"의 (?:광역 )?(?:추가 |연쇄 )?(?:물리|특수)\s*피해", effect)]
            ratios = [{"Values": [float(v) / 100 for v in match.groups()]}
                      for match in re.finditer(TRIPLE + "%", effect)]
            rows.append(dict(Dex=dex, Id=ids[dex], Name=row[1], Cost=row[2], Skill=row[3],
                             Effect=effect, Notes=row[6] or "", Damage=damage, Ratios=ratios,
                             HP=list(stat[3:6]), Attack=list(stat[6:9]), AttackSpeed=stat[9],
                             MaxEnergy=stat[10], StartingEnergy=stat[11], Armor=stat[12],
                             MagicResistance=stat[13], AttackRange=stat[14], Role=stat[15]))
        if len(rows) != 109 or len({row["Id"] for row in rows}) != 109:
            raise ValueError("Expected 109 unique Pokemon")
        # This importer supports this schema and formula; future spreadsheet revisions need review.
        if "최종 주문력 ÷ 100" not in str(workbook["피해량 기준"]["B12"].value):
            raise ValueError("Unsupported damage formula")
        return dict(Source=Path(path).name, Formula="BaseDamage * SpellPower / 100",
                    Rows=rows, Families=baseline["Families"])
    finally:
        workbook.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    plan = extract(args.workbook)
    output = CONTENT / "SkillPlan.json"
    output.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Extracted {len(plan['Rows'])} Pokemon and {sum(len(r['Damage']) for r in plan['Rows'])} damage triplets")
